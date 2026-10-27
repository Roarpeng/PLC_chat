"""IO 表导入:PLCTags.xlsx / CSV → 画布标签(契约 24 · A 线)。

TIA Portal「PLC 变量表」导出的 ``PLCTags.xlsx``(列 = Name / Path /
Data Type / Logical Address / Comment …)或同列名 CSV 文本,解析为统一
entries,再折算成 ``CanvasTag(status="draft")`` 合并进 ``canvas.tags``。
最小破坏原则:已有同名标签一律不覆盖(保持原 status/地址/注释),仅新增
缺失项;blocks/HMI 完全不动。
"""

from __future__ import annotations

import csv
import io
import os
from pathlib import Path
from typing import Any

from plc_engine.canvas_model import CanvasState, CanvasTag

XLSX_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIX = ".csv"

# 表头别名(大小写/空白不敏感,兼容中英文表头)。key 归一化后匹配,
# value 为统一 entry 字段;path 仅透传保留在 entry 里,不进 CanvasTag。
_HEADER_ALIASES: dict[str, str] = {
    "name": "name",
    "tag name": "name",
    "标签": "name",
    "变量名": "name",
    "path": "path",
    "路径": "path",
    "data type": "data_type",
    "datatype": "data_type",
    "type": "data_type",
    "数据类型": "data_type",
    "logical address": "logical_address",
    "logicaladdress": "logical_address",
    "address": "logical_address",
    "logicaladdress(absolute)": "logical_address",
    "地址": "logical_address",
    "comment": "comment",
    "注释": "comment",
}

# 识别为表头行的最低条件:name 列 + (data_type 或 logical_address) 列。
_HEADER_MAX_SCAN_ROWS = 5


class IoImportRejected(ValueError):
    """Import rejected before any canvas mutation; the router maps it to HTTP 400."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _norm_header(value: Any) -> str:
    """Normalize a header cell: str -> strip -> lowercase -> collapse whitespace."""
    return " ".join(str(value or "").strip().lower().split())


def _cell_str(value: Any) -> str:
    """Normalize a data cell to a clean string (numbers, None, multiline text)."""
    if value is None:
        return ""
    return str(value).strip()


def _map_header_columns(headers: list[Any]) -> dict[int, str] | None:
    """Map column index -> entry field; None when the row is not a PLCTags header."""
    mapping: dict[int, str] = {}
    for idx, raw in enumerate(headers):
        field = _HEADER_ALIASES.get(_norm_header(raw))
        if field and field not in mapping.values():
            mapping[idx] = field
    if "name" not in mapping.values():
        return None
    if "data_type" not in mapping.values() and "logical_address" not in mapping.values():
        return None
    return mapping


def _row_to_entry(row: tuple[Any, ...], mapping: dict[int, str]) -> dict[str, str] | None:
    """Convert one spreadsheet/CSV row to a unified entry; None when the row is junk."""
    if _map_header_columns(list(row)) is not None:
        return None  # 表头行(或多页导出时重复出现的表头)不当作数据
    entry: dict[str, str] = {}
    for idx, field in mapping.items():
        raw = row[idx] if idx < len(row) else None
        entry[field] = _cell_str(raw)
    name = entry.get("name") or ""
    if not name:
        return None  # 空行/小计行——直接丢弃
    entry.setdefault("data_type", "")
    entry.setdefault("logical_address", "")
    entry.setdefault("comment", "")
    return entry


def parse_plctags_xlsx(path: str | Path) -> list[dict[str, str]]:
    """Parse a TIA-exported ``PLCTags.xlsx`` into unified entries.

    列名映射 Name / Data Type / Logical Address / Comment(Path 透传保留);
    自动跳过空 Name 行。表头不可识别或文件损坏时抛 ``IoImportRejected``
    (router 映射为 HTTP 400)。
    """
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from zipfile import BadZipFile

    target = Path(path)
    try:
        workbook = load_workbook(target, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError) as exc:
        raise IoImportRejected(
            "IO_XLSX_UNREADABLE",
            f"无法读取 xlsx 文件 {target.name}: {exc}",
        ) from exc

    entries: list[dict[str, str]] = []
    try:
        for worksheet in workbook.worksheets:
            # IO 表规模在百行级,一次性物化后定位表头,避免二次 iter_rows
            # 把表头行当数据行吞进去(read_only 模式流不可回退)。
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            mapping: dict[int, str] | None = None
            data_start = 0
            for offset, row in enumerate(rows[:_HEADER_MAX_SCAN_ROWS]):
                mapping = _map_header_columns(row)
                if mapping is not None:
                    data_start = offset + 1
                    break
            if mapping is None:
                continue  # 该 sheet 没有 PLCTags 表头(如说明页),看下一个
            for row in rows[data_start:]:
                entry = _row_to_entry(tuple(row), mapping)
                if entry is not None:
                    entries.append(entry)
            break  # 只取第一个命中表头的 sheet(TIA 导出即单表 "PLC Tags")
    finally:
        workbook.close()

    if not entries:
        # 走到这里说明没有任何 sheet 带可识别表头(有表头但零数据行也归此码)。
        raise IoImportRejected(
            "IO_TABLE_HEADER_MISSING",
            "xlsx 中未找到 PLCTags 表头(需 Name + Data Type/Logical Address 列)",
        )
    return entries


def _decode_csv_bytes(data: bytes) -> str:
    """Decode CSV bytes: utf-8(BOM) -> gbk(TIA 中文 Windows 常见) -> latin-1 兜底."""
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _csv_rows(text: str) -> list[list[str]]:
    """Split CSV text into rows with sniffed delimiter (, ; or tab)."""
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel  # type: ignore[assignment]
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def parse_plctags_csv(source: str | Path) -> list[dict[str, str]]:
    """Parse same-header CSV — accepts a file path or raw CSV text.

    列名与 xlsx 相同(Name/Data Type/Logical Address/Comment…);空 Name 行
    丢弃。表头不可识别时抛 ``IoImportRejected``。
    """
    if isinstance(source, Path):
        text = _decode_csv_bytes(source.read_bytes())
    elif "\n" not in source and "\r" not in source and os.path.exists(source):
        text = _decode_csv_bytes(Path(source).read_bytes())
    else:
        text = source

    rows = _csv_rows(text)
    mapping: dict[int, str] | None = None
    for offset, row in enumerate(rows[:_HEADER_MAX_SCAN_ROWS]):
        mapping = _map_header_columns(row)
        if mapping is not None:
            rows = rows[offset + 1 :]
            break
    if mapping is None:
        raise IoImportRejected(
            "IO_TABLE_HEADER_MISSING",
            "CSV 中未找到 PLCTags 表头(需 Name + Data Type/Logical Address 列)",
        )

    entries: list[dict[str, str]] = []
    for row in rows:
        entry = _row_to_entry(tuple(row), mapping)
        if entry is not None:
            entries.append(entry)
    if not entries:
        raise IoImportRejected(
            "IO_NO_ENTRIES",
            "CSV 解析后没有有效标签行(Name 列全为空)",
        )
    return entries


def io_entries_to_canvas_tags(entries: list[dict[str, Any]]) -> dict[str, CanvasTag]:
    """Unified entries -> ``{name: CanvasTag}`` with status="draft".

    draft 状态是刻意选择:导入标签尚未写回 TIA,只有 draft 标签会触发
    标签表整表重写(与块导入的最小破坏原则一致)。
    """
    tags: dict[str, CanvasTag] = {}
    for entry in entries:
        name = _cell_str(entry.get("name"))
        if not name:
            continue
        tags[name] = CanvasTag(
            name=name,
            data_type=_cell_str(entry.get("data_type")),
            logical_address=_cell_str(entry.get("logical_address")),
            comment=_cell_str(entry.get("comment")),
            status="draft",
        )
    return tags


def import_io_table(
    source: str | Path | None = None,
    canvas: CanvasState | None = None,
    project_name: str = "",
    *,
    csv_text: str = "",
) -> dict[str, Any]:
    """Import an IO table (xlsx/csv path or CSV text) and merge tags into a canvas.

    合并语义(最小破坏):同名标签不覆盖——已有标签的 status/地址/注释保持
    原样,仅补齐缺失项;``blocks`` / ``hmi_widgets`` 完全不动。返回
    ``{tags_created, tags_skipped, tags_total, canvas}``。
    """
    if canvas is None:
        canvas = CanvasState(project_name=project_name.strip())
    elif project_name.strip():
        canvas.project_name = project_name.strip()

    if csv_text.strip():
        entries = parse_plctags_csv(csv_text)
    elif source is not None and str(source).strip():
        path = Path(str(source))
        suffix = path.suffix.lower()
        if suffix in XLSX_SUFFIXES:
            entries = parse_plctags_xlsx(path)
        elif suffix == CSV_SUFFIX:
            entries = parse_plctags_csv(path)
        else:
            raise IoImportRejected(
                "IO_SOURCE_UNSUPPORTED",
                f"不支持的 IO 表格式 '{suffix or '(无后缀)'}',仅支持 .xlsx/.xlsm/.csv",
            )
    else:
        raise IoImportRejected(
            "IO_SOURCE_REQUIRED", "file_path 或 file_content(CSV 文本)至少提供一个"
        )

    tags_created = 0
    tags_skipped = 0
    for name, tag in io_entries_to_canvas_tags(entries).items():
        if name in canvas.tags:
            tags_skipped += 1  # 已有标签(含 synced)不覆盖
            continue
        canvas.tags[name] = tag
        tags_created += 1

    return {
        "tags_created": tags_created,
        "tags_skipped": tags_skipped,
        "tags_total": len(canvas.tags),
        "entries_parsed": len(entries),
        "canvas": canvas.to_dict(),
    }
